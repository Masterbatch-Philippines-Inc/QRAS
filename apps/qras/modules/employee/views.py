from django.shortcuts import render, redirect
from apps.qras.modules.auth.decorators import permission_required
from apps.qras.modules.auth.helpers import get_dept_queryset_filter, get_self_exclude
from django.contrib.auth.mixins import LoginRequiredMixin
from django.utils.decorators import method_decorator
from django.http import JsonResponse, HttpResponse
from django.views import View
from django.shortcuts import get_object_or_404
from django.db.models import Q, Count, Prefetch
from datetime import date
import json
import openpyxl
from io import BytesIO

from apps.qras.models.employee import (
    Employee,
    EmployeeContact, 
    EmployeeGovernmentID,
    EmployeeDetails, 
    EmployeeGroup,
    EmployeeGroupMembership
)
from apps.qras.models.access import (
    Department,
    Position
)

# ─────────────────────────────────────────
#   SHARED HELPERS
# ─────────────────────────────────────────

def clean_date(val):
    from datetime import datetime
    v = (val or '').strip()
    if not v or '_' in v:
        return None
    for fmt in ('%m/%d/%Y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(v, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return None

def clean_masked(val):
    v = (val or '').strip()
    return '' if not v or '_' in v else v

def clean_phone(val):
    v = (val or '').strip()
    if not v or '_' in v:
        return ''
    digits = ''.join(c for c in v if c.isdigit())
    return '0' + digits[-10:] if len(digits) >= 10 else digits

# ─────────────────────────────────────────
#   EMPLOYEE VIEWS
# ─────────────────────────────────────────

@method_decorator([permission_required('employees', 'read')], name='dispatch')
class EmployeeListView(LoginRequiredMixin, View):
    template_name = 'pages/employee/records.django'

    def get(self, request):
        dept_filter = get_dept_queryset_filter(request, dept_field_path='department')
        if dept_filter is None:
            return render(request, self.template_name, {'dept_warning': True})

        self_exclude = get_self_exclude(request)

        employees = Employee.objects.filter(**dept_filter).select_related(
            'department', 'position', 'contact', 'gov_ids', 'details'
        ).order_by('last_name', 'first_name')
        if self_exclude:
            employees = employees.exclude(pk=self_exclude)

        departments = Department.objects.order_by('name')
        positions   = Position.objects.select_related('department').order_by('title')
        positions_json = json.dumps([
            {'id': p.id, 'title': p.title, 'department_id': p.department_id}
            for p in positions
        ])

        return render(request, self.template_name, {
            'employees':   employees,
            'departments': departments,
            'positions':   positions,
            'positions_json':  positions_json,
        })


@method_decorator([permission_required('employees', 'create')], name='dispatch')
class EmployeeCreateView(LoginRequiredMixin, View):
    template_name = 'pages/employee/single_add.django'

    def get(self, request):
        positions = Position.objects.filter(is_active=True).select_related('department').order_by('title')
        return render(request, self.template_name, {
            'departments':     Department.objects.filter(is_active=True).order_by('name'),
            'positions':       positions,
            'positions_json':  json.dumps([
                {'id': p.id, 'title': p.title, 'department_id': p.department_id}
                for p in positions
            ]),
            'employee_groups': EmployeeGroup.objects.filter(is_active=True).order_by('group_name'),
        })

    def post(self, request):
        first_name          = request.POST.get('first_name', '').strip()
        middle_name         = request.POST.get('middle_name', '').strip()
        last_name           = request.POST.get('last_name', '').strip()
        suffix_name         = request.POST.get('suffix_name', '').strip()
        emp_id_raw          = request.POST.get('emp_id', '').replace(' ', '').strip()
        emp_id              = emp_id_raw if emp_id_raw and '_' not in emp_id_raw else ''
        department_id       = request.POST.get('department_id', '').strip()
        position_id         = request.POST.get('position_id', '').strip()
        group_id            = request.POST.get('group_id', '').strip()
        hire_date           = clean_date(request.POST.get('hire_date'))
        date_regularized    = clean_date(request.POST.get('date_regularized'))
        birth_date          = clean_date(request.POST.get('birth_date'))
        contact_no          = clean_phone(request.POST.get('contact_no'))
        address_1           = request.POST.get('address_1', '').strip()
        address_2           = request.POST.get('address_2', '').strip()
        sos_contact_person  = request.POST.get('sos_contact_person', '').strip()
        sos_contact_no      = clean_phone(request.POST.get('sos_contact_no'))
        tin_no              = clean_masked(request.POST.get('tin_no'))
        sss_no              = clean_masked(request.POST.get('sss_no'))
        philhealth_no       = clean_masked(request.POST.get('philhealth_no'))
        pagibig_no          = clean_masked(request.POST.get('pagibig_no'))

        errors = {}
        if not first_name: errors['first_name'] = 'First name is required.'
        if not last_name:  errors['last_name']  = 'Last name is required.'
        if not emp_id:
            errors['emp_id'] = 'Employee ID is required.'
        elif not emp_id.isdigit():
            errors['emp_id'] = 'Employee ID must be a number.'
        elif Employee.objects.filter(employee_id=emp_id).exists():
            errors['emp_id'] = 'Employee ID already exists.'

        department = None
        if department_id:
            try:
                department = Department.objects.get(pk=department_id)
            except Department.DoesNotExist:
                errors['department_id'] = 'Selected department does not exist.'

        position = None
        if position_id:
            try:
                position = Position.objects.get(pk=position_id)
            except Position.DoesNotExist:
                errors['position_id'] = 'Selected position does not exist.'

        if errors:
            return JsonResponse({'success': False, 'errors': errors}, status=400)

        emp = Employee.objects.create(
            first_name=first_name,
            middle_name=middle_name,
            last_name=last_name,
            suffix_name=suffix_name,
            employee_id=emp_id,
            department=department,
            position=position,
            hire_date=hire_date,
            date_regularized=date_regularized,
        )

        if birth_date:
            EmployeeDetails.objects.create(employee=emp, birth_date=birth_date)

        if any([contact_no, address_1, address_2, sos_contact_person, sos_contact_no]):
            EmployeeContact.objects.create(
                employee=emp,
                contact_no=contact_no,
                address_1=address_1,
                address_2=address_2,
                sos_contact_person=sos_contact_person,
                sos_contact_no=sos_contact_no,
            )

        if any([tin_no, sss_no, philhealth_no, pagibig_no]):
            EmployeeGovernmentID.objects.create(
                employee=emp,
                tin_no=tin_no,
                sss_no=sss_no,
                philhealth_no=philhealth_no,
                pagibig_no=pagibig_no,
            )

        if group_id:
            try:
                group = EmployeeGroup.objects.get(pk=group_id)
                EmployeeGroupMembership.objects.get_or_create(employee=emp, group=group)
            except EmployeeGroup.DoesNotExist:
                pass

        return JsonResponse({
            'success': True,
            'message': f'Employee "{emp.last_name}, {emp.first_name}" has been added.',
        })


@method_decorator([permission_required('employees', 'update')], name='dispatch')
class EmployeeUpdateView(LoginRequiredMixin, View):
    def post(self, request, pk):
        emp = get_object_or_404(Employee, pk=pk)

        # ── Core fields ──
        first_name       = request.POST.get('first_name', '').strip()
        middle_name      = request.POST.get('middle_name', '').strip()
        last_name        = request.POST.get('last_name', '').strip()
        suffix_name      = request.POST.get('suffix_name', '').strip()
        department_id    = request.POST.get('department_id', '').strip()
        position_id      = request.POST.get('position_id', '').strip()
        hire_date        = clean_date(request.POST.get('hire_date'))
        date_regularized = clean_date(request.POST.get('date_regularized'))

        # ── Contact fields ──
        contact_no         = clean_phone(request.POST.get('contact_no'))
        address_1          = request.POST.get('address_1', '').strip()
        address_2          = request.POST.get('address_2', '').strip()
        sos_contact_person = request.POST.get('sos_contact_person', '').strip()
        sos_contact_no     = clean_phone(request.POST.get('sos_contact_no'))

        # ── Details fields ──
        birth_date = clean_date(request.POST.get('birth_date'))

        # ── Government ID fields ──
        tin_no        = clean_masked(request.POST.get('tin_no'))
        sss_no        = clean_masked(request.POST.get('sss_no'))
        philhealth_no = clean_masked(request.POST.get('philhealth_no'))
        pagibig_no    = clean_masked(request.POST.get('pagibig_no'))

        errors = {}
        if not first_name: errors['first_name'] = 'First name is required.'
        if not last_name:  errors['last_name']  = 'Last name is required.'

        department = None
        if department_id:
            try:
                department = Department.objects.get(pk=department_id)
            except Department.DoesNotExist:
                errors['department_id'] = 'Selected department does not exist.'

        position = None
        if position_id:
            try:
                position = Position.objects.get(pk=position_id)
            except Position.DoesNotExist:
                errors['position_id'] = 'Selected position does not exist.'

        if errors:
            return JsonResponse({'success': False, 'errors': errors}, status=400)

        # ── Save Employee ──
        emp.first_name       = first_name
        emp.middle_name      = middle_name
        emp.last_name        = last_name
        emp.suffix_name      = suffix_name
        emp.department       = department
        emp.position         = position
        emp.hire_date        = hire_date
        emp.date_regularized = date_regularized
        emp.save()

        # ── Save Contact (create if not exists) ──
        contact, _ = EmployeeContact.objects.get_or_create(employee=emp)
        contact.contact_no         = contact_no
        contact.address_1          = address_1
        contact.address_2          = address_2
        contact.sos_contact_person = sos_contact_person
        contact.sos_contact_no     = sos_contact_no
        contact.save()

        # ── Save Details (create if not exists) ──
        details, _ = EmployeeDetails.objects.get_or_create(employee=emp)
        details.birth_date = birth_date
        details.save()

        # ── Save Government IDs (create if not exists) ──
        gov, _ = EmployeeGovernmentID.objects.get_or_create(employee=emp)
        gov.tin_no        = tin_no
        gov.sss_no        = sss_no
        gov.philhealth_no = philhealth_no
        gov.pagibig_no    = pagibig_no
        gov.save()

        return JsonResponse({
            'success': True,
            'message': f'Employee "{emp.first_name} {emp.last_name}" updated successfully.',
            'employee': {
                'id':         emp.pk,
                'name':       f'{emp.last_name}, {emp.first_name}',
                'department': emp.department.name if emp.department else '—',
                'position':   emp.position.title  if emp.position  else '—',
            }
        })


@method_decorator([permission_required('employees', 'update')], name='dispatch')
class EmployeeBatchDeactivateView(LoginRequiredMixin, View):
    def post(self, request):
        # Admin-only guard
        if not request.user.role or request.user.role.name != 'Administrator':
            return JsonResponse({'success': False, 'message': 'Only Administrators can perform this action.'}, status=403)

        try:
            data = json.loads(request.body)
            emp_ids = data.get('employee_ids', [])
            action  = data.get('action', 'deactivate')  # 'deactivate' or 'reactivate'
        except json.JSONDecodeError:
            return JsonResponse({'success': False, 'message': 'Invalid request body.'}, status=400)

        if not emp_ids:
            return JsonResponse({'success': False, 'message': 'No employees selected.'}, status=400)

        if action == 'reactivate':
            employees = Employee.objects.filter(pk__in=emp_ids, is_active=False, is_resigned=False)
            count = employees.count()
            employees.update(is_active=True)
            return JsonResponse({
                'success': True,
                'count': count,
                'message': f'{count} employee(s) restored to active.',
            })

        # default: deactivate
        employees = Employee.objects.filter(pk__in=emp_ids, is_active=True)
        count = employees.count()

        employees.update(is_active=False)

        from apps.qras.models.employee import EmployeeSchedule
        EmployeeSchedule.objects.filter(employee_id__in=emp_ids, is_active=True).update(is_active=False)

        return JsonResponse({
            'success': True,
            'count': count,
            'message': f'{count} employee(s) set to inactive.',
        })


@method_decorator([permission_required('employees', 'update')], name='dispatch')
class EmployeeResignView(LoginRequiredMixin, View):
    def post(self, request, pk):
        emp = get_object_or_404(Employee, pk=pk)

        if emp.is_resigned:
            emp.is_resigned   = False
            emp.is_active     = True
            emp.date_resigned = None
        else:
            emp.is_resigned   = True
            emp.is_active     = False
            emp.date_resigned = date.today()

        emp.save(update_fields=['is_resigned', 'is_active', 'date_resigned'])

        action = 'restored' if emp.is_active else 'marked as resigned'
        return JsonResponse({
            'success':     True,
            'is_resigned': emp.is_resigned,
            'is_active':   emp.is_active,
            'message':     f'Employee "{emp.first_name} {emp.last_name}" {action}.',
        })


# ─────────────────────────────────────────
#   DEPARTMENT VIEWS
# ─────────────────────────────────────────

@method_decorator([permission_required('settings_departments', 'read')], name='dispatch')
class DepartmentListView(LoginRequiredMixin, View):
    template_name = "pages/settings/departments.django"

    def get(self, request):
        search        = request.GET.get("q", "").strip()
        show_inactive = request.GET.get("show_inactive", "false") == "true"
        qs = Department.objects.annotate(employee_count=Count("employee"))
        if hasattr(Department, "is_active"):
            if not show_inactive:
                qs = qs.filter(is_active=True)
        if search:
            qs = qs.filter(Q(name__icontains=search) | Q(code__icontains=search))
        return render(request, self.template_name, {
            "departments":   qs.order_by("name"),
            "search":        search,
            "show_inactive": show_inactive,
        })


@method_decorator([permission_required('settings_departments', 'create')], name='dispatch')
class DepartmentCreateView(LoginRequiredMixin, View):
    def post(self, request):
        name = request.POST.get("name", "").strip()
        code = request.POST.get("code", "").strip()
        errors = {}
        if not name: errors["name"] = "Department name is required."
        if not code: errors["code"] = "Department code is required."
        if Department.objects.filter(code__iexact=code).exists():
            errors["code"] = "A department with this code already exists."
        if errors:
            return JsonResponse({"success": False, "errors": errors}, status=400)
        dept = Department.objects.create(name=name, code=code.upper())
        return JsonResponse({"success": True, "message": f'Department "{dept.name}" created.', "department": {"id": dept.id, "name": dept.name, "code": dept.code}})


@method_decorator([permission_required('settings_departments', 'update')], name='dispatch')
class DepartmentUpdateView(LoginRequiredMixin, View):
    def get(self, request, pk):
        dept = get_object_or_404(Department, pk=pk)
        return JsonResponse({"id": dept.id, "name": dept.name, "code": dept.code})

    def post(self, request, pk):
        dept = get_object_or_404(Department, pk=pk)
        name = request.POST.get("name", "").strip()
        code = request.POST.get("code", "").strip()
        errors = {}
        if not name: errors["name"] = "Department name is required."
        if not code: errors["code"] = "Department code is required."
        if Department.objects.filter(code__iexact=code).exclude(pk=pk).exists():
            errors["code"] = "A department with this code already exists."
        if errors:
            return JsonResponse({"success": False, "errors": errors}, status=400)
        dept.name = name
        dept.code = code.upper()
        dept.save()
        return JsonResponse({"success": True, "message": f'Department "{dept.name}" updated.', "department": {"id": dept.id, "name": dept.name, "code": dept.code}})


@method_decorator([permission_required('settings_departments', 'delete')], name='dispatch')
class DepartmentSoftDeleteView(LoginRequiredMixin, View):
    def post(self, request, pk):
        dept = get_object_or_404(Department, pk=pk)
        if not hasattr(dept, "is_active"):
            return JsonResponse({"success": False, "message": "Soft delete not configured."}, status=501)
        dept.is_active = not dept.is_active
        dept.save(update_fields=["is_active"])
        action = "restored" if dept.is_active else "deactivated"
        return JsonResponse({"success": True, "is_active": dept.is_active, "message": f'Department "{dept.name}" {action}.'})
    
@method_decorator([permission_required('settings_departments', 'read')], name='dispatch')
class DepartmentEmployeesView(LoginRequiredMixin, View):
    def get(self, request, pk):
        employees = Employee.objects.filter(
            department_id=pk, is_active=True, is_resigned=False
        ).select_related('position').order_by('last_name', 'first_name')
        data = [
            {
                'name':     f"{e.last_name}, {e.first_name}",
                'position': e.position.title if e.position else '—',
            }
            for e in employees
        ]
        return JsonResponse({'employees': data})


# ─────────────────────────────────────────
#   POSITION VIEWS
# ─────────────────────────────────────────

@method_decorator([permission_required('settings_positions', 'read')], name='dispatch')
class PositionListView(LoginRequiredMixin, View):
    template_name = "pages/settings/positions.django"

    def get(self, request):
        search        = request.GET.get("q", "").strip()
        dept_filter   = request.GET.get("department", "")
        show_inactive = request.GET.get("show_inactive", "false") == "true"
        qs = Position.objects.select_related("department").annotate(employee_count=Count("employee"))
        if hasattr(Position, "is_active"):
            if not show_inactive:
                qs = qs.filter(is_active=True)
        if search:
            qs = qs.filter(Q(title__icontains=search) | Q(department__name__icontains=search))
        if dept_filter:
            qs = qs.filter(department_id=dept_filter)
        departments = Department.objects.order_by("name")
        if hasattr(Department, "is_active"):
            departments = departments.filter(is_active=True)
        return render(request, self.template_name, {
            "positions":    qs.order_by("department__name", "title"),
            "departments":  departments,
            "search":       search,
            "dept_filter":  dept_filter,
            "show_inactive": show_inactive,
        })


@method_decorator([permission_required('settings_positions', 'create')], name='dispatch')
class PositionCreateView(LoginRequiredMixin, View):
    def post(self, request):
        title         = request.POST.get("title", "").strip()
        department_id = request.POST.get("department_id", "").strip()
        errors = {}
        if not title: errors["title"] = "Position title is required."
        department = None
        if department_id:
            try:
                department = Department.objects.get(pk=department_id)
            except Department.DoesNotExist:
                errors["department_id"] = "Selected department does not exist."
        if errors:
            return JsonResponse({"success": False, "errors": errors}, status=400)
        position = Position.objects.create(title=title, department=department)
        return JsonResponse({"success": True, "message": f'Position "{position.title}" created.', "position": {"id": position.id, "title": position.title, "department": position.department.name if position.department else "—", "department_id": position.department_id}})


@method_decorator([permission_required('settings_positions', 'update')], name='dispatch')
class PositionUpdateView(LoginRequiredMixin, View):
    def get(self, request, pk):
        position = get_object_or_404(Position, pk=pk)
        return JsonResponse({"id": position.id, "title": position.title, "department_id": position.department_id or ""})

    def post(self, request, pk):
        position      = get_object_or_404(Position, pk=pk)
        title         = request.POST.get("title", "").strip()
        department_id = request.POST.get("department_id", "").strip()
        errors = {}
        if not title: errors["title"] = "Position title is required."
        department = None
        if department_id:
            try:
                department = Department.objects.get(pk=department_id)
            except Department.DoesNotExist:
                errors["department_id"] = "Selected department does not exist."
        if errors:
            return JsonResponse({"success": False, "errors": errors}, status=400)
        position.title      = title
        position.department = department
        position.save()
        return JsonResponse({"success": True, "message": f'Position "{position.title}" updated.', "position": {"id": position.id, "title": position.title, "department": position.department.name if position.department else "—", "department_id": position.department_id}})


@method_decorator([permission_required('settings_positions', 'delete')], name='dispatch')
class PositionSoftDeleteView(LoginRequiredMixin, View):
    def post(self, request, pk):
        position = get_object_or_404(Position, pk=pk)
        if not hasattr(position, "is_active"):
            return JsonResponse({"success": False, "message": "Soft delete not configured."}, status=501)
        position.is_active = not position.is_active
        position.save(update_fields=["is_active"])
        action = "restored" if position.is_active else "deactivated"
        return JsonResponse({"success": True, "is_active": position.is_active, "message": f'Position "{position.title}" {action}.'})

@method_decorator([permission_required('settings_positions', 'read')], name='dispatch')
class PositionEmployeesView(LoginRequiredMixin, View):
    def get(self, request, pk):
        employees = Employee.objects.filter(
            position_id=pk, is_active=True, is_resigned=False
        ).select_related('department').order_by('last_name', 'first_name')
        data = [
            {
                'name':       f"{e.last_name}, {e.first_name}",
                'department': e.department.name if e.department else '—',
            }
            for e in employees
        ]
        return JsonResponse({'employees': data})

# ─────────────────────────────────────────
#   EMPLOYEE GROUP VIEWS
# ─────────────────────────────────────────

@method_decorator([permission_required('settings_employee_groups', 'update')], name='dispatch')
class EmployeeGroupListView(LoginRequiredMixin, View):
    template_name = 'pages/settings/employee_group.django'

    def get(self, request):
        search        = request.GET.get('q', '').strip()
        show_inactive = request.GET.get('show_inactive', 'false') == 'true'

        qs = EmployeeGroup.objects.annotate(member_count=Count('memberships'))

        if not show_inactive:
            qs = qs.filter(is_active=True)

        if search:
            qs = qs.filter(group_name__icontains=search)

        # Pass employees with their current group memberships
        dept_filter = get_dept_queryset_filter(request, dept_field_path='department')
        if dept_filter is None:
            return render(request, self.template_name, {'dept_warning': True})

        self_exclude = get_self_exclude(request)

        employees = Employee.objects.filter(
            is_active=True, is_resigned=False, **dept_filter
        ).prefetch_related(
            Prefetch(
                'group_memberships',
                queryset=EmployeeGroupMembership.objects.select_related('group'),
                to_attr='current_memberships'
            )
        ).order_by('last_name', 'first_name')
        if self_exclude:
            employees = employees.exclude(pk=self_exclude)

        active_groups = EmployeeGroup.objects.filter(is_active=True).order_by('group_name')

        return render(request, self.template_name, {
            'groups':        qs.order_by('group_name'),
            'search':        search,
            'show_inactive': show_inactive,
            'employees':     employees,
            'active_groups': active_groups,
        })


@method_decorator([permission_required('settings_employee_groups', 'create')], name='dispatch')
class EmployeeGroupCreateView(LoginRequiredMixin, View):
    def post(self, request):
        group_name         = request.POST.get('group_name', '').strip()
        ot_filing_exempted = request.POST.get('ot_filing_exempted') == 'true'
 
        errors = {}
        if not group_name:
            errors['group_name'] = 'Group name is required.'
        if EmployeeGroup.objects.filter(group_name__iexact=group_name).exists():
            errors['group_name'] = 'A group with this name already exists.'
 
        if errors:
            return JsonResponse({'success': False, 'errors': errors}, status=400)
 
        group = EmployeeGroup.objects.create(
            group_name=group_name,
            ot_filing_exempted=ot_filing_exempted,
        )
 
        return JsonResponse({
            'success': True,
            'message': f'Group "{group.group_name}" created.',
            'group': {
                'id':                 group.id,
                'group_name':         group.group_name,
                'ot_filing_exempted': group.ot_filing_exempted,
            }
        })


@method_decorator([permission_required('settings_employee_groups', 'update')], name='dispatch')
class EmployeeGroupUpdateView(LoginRequiredMixin, View):
    def get(self, request, pk):
        group = get_object_or_404(EmployeeGroup, pk=pk)
        return JsonResponse({
            'id':                 group.id,
            'group_name':         group.group_name,
            'ot_filing_exempted': group.ot_filing_exempted,
        })
 
    def post(self, request, pk):
        group              = get_object_or_404(EmployeeGroup, pk=pk)
        group_name         = request.POST.get('group_name', '').strip()
        ot_filing_exempted = request.POST.get('ot_filing_exempted') == 'true'
 
        errors = {}
        if not group_name:
            errors['group_name'] = 'Group name is required.'
        if EmployeeGroup.objects.filter(group_name__iexact=group_name).exclude(pk=pk).exists():
            errors['group_name'] = 'A group with this name already exists.'
 
        if errors:
            return JsonResponse({'success': False, 'errors': errors}, status=400)
 
        group.group_name         = group_name
        group.ot_filing_exempted = ot_filing_exempted
        group.save()
 
        return JsonResponse({
            'success': True,
            'message': f'Group "{group.group_name}" updated.',
            'group': {
                'id':                 group.id,
                'group_name':         group.group_name,
                'ot_filing_exempted': group.ot_filing_exempted,
            }
        })


@method_decorator([permission_required('settings_employee_groups', 'delete')], name='dispatch')
class EmployeeGroupSoftDeleteView(LoginRequiredMixin, View):
    def post(self, request, pk):
        group          = get_object_or_404(EmployeeGroup, pk=pk)
        group.is_active = not group.is_active
        group.save(update_fields=['is_active'])
 
        action = 'restored' if group.is_active else 'deactivated'
        return JsonResponse({
            'success':   True,
            'is_active': group.is_active,
            'message':   f'Group "{group.group_name}" {action}.',
        })


@method_decorator([permission_required('settings_employee_groups', 'update')], name='dispatch')
class EmployeeGroupAssignView(LoginRequiredMixin, View):
    def post(self, request):
        try:
            data      = json.loads(request.body)
            group_id  = data.get('group_id')
            emp_ids   = data.get('employee_ids', [])
 
            group     = EmployeeGroup.objects.get(id=group_id)
            employees = Employee.objects.filter(employee_id__in=emp_ids)
 
            for emp in employees:
                EmployeeGroupMembership.objects.get_or_create(
                    employee=emp,
                    group=group,
                )
 
            return JsonResponse({
                'status':     'assigned',
                'group_name': group.group_name,
                'updated':    employees.count(),
            })
 
        except EmployeeGroup.DoesNotExist:
            return JsonResponse({'error': 'Group not found.'}, status=404)
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=400)


@method_decorator([permission_required('settings_employee_groups', 'update')], name='dispatch')
class EmployeeGroupUnassignView(LoginRequiredMixin, View):
    def get(self, request):
        """Return employee + group info before confirming unassign."""
        emp_id   = request.GET.get('employee_id')
        group_id = request.GET.get('group_id')
 
        try:
            membership = EmployeeGroupMembership.objects.select_related(
                'employee', 'group'
            ).get(employee__employee_id=emp_id, group_id=group_id)
 
            return JsonResponse({
                'employee_id':   membership.employee.employee_id,
                'employee_name': f'{membership.employee.last_name}, {membership.employee.first_name}',
                'group_name':    membership.group.group_name,
                'group_id':      membership.group.id,
            })
 
        except EmployeeGroupMembership.DoesNotExist:
            return JsonResponse({'error': 'Membership not found.'}, status=404)
 
    def post(self, request):
        """Remove employee from group."""
        try:
            data     = json.loads(request.body)
            emp_id   = data.get('employee_id')
            group_id = data.get('group_id')
 
            membership = EmployeeGroupMembership.objects.get(
                employee__employee_id=emp_id,
                group_id=group_id,
            )
            group_name = membership.group.group_name
            membership.delete()
 
            return JsonResponse({
                'status':     'unassigned',
                'message':    f'Employee removed from "{group_name}".',
            })
 
        except EmployeeGroupMembership.DoesNotExist:
            return JsonResponse({'error': 'Membership not found.'}, status=404)
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=400)

# ─────────────────────────────────────────
#   EMPLOYEE BATCH IMPORT VIEWS
# ─────────────────────────────────────────

@method_decorator([permission_required('employees', 'create')], name='dispatch')
class EmployeeBatchUploadView(LoginRequiredMixin, View):
    template_name = 'pages/employee/batch_add.django'

    def get(self, request):
        return render(request, self.template_name, {
            'departments': Department.objects.filter(is_active=True).order_by('name'),
            'positions':   Position.objects.filter(is_active=True).select_related('department').order_by('title'),
        })


@method_decorator([permission_required('employees', 'create')], name='dispatch')
class EmployeeBatchPreviewView(LoginRequiredMixin, View):

    def post(self, request):
        file = request.FILES.get('excel_file')
        if not file:
            return JsonResponse({'error': 'No file provided.'}, status=400)
        if not file.name.lower().endswith('.xlsx'):
            return JsonResponse({'error': 'Only .xlsx files are supported. Please use the provided template.'}, status=400)

        try:
            wb = openpyxl.load_workbook(file, data_only=True)
            ws = wb.active
        except Exception as e:
            return JsonResponse({'error': f'Could not read file: {e}'}, status=400)

        # Maps any known header variant (lowercased+stripped) → canonical field name
        HEADER_ALIASES = {
            'emp. #':               'employee_id',
            'emp #':                'employee_id',
            'emp_id':               'employee_id',
            'employee id':          'employee_id',
            'employee_id':          'employee_id',
            'first name':           'first_name',
            'first_name':           'first_name',
            'last name':            'last_name',
            'last_name':            'last_name',
            'middle name':          'middle_name',
            'middle_name':          'middle_name',
            'suffix':               'suffix_name',
            'suffix name':          'suffix_name',
            'suffix_name':          'suffix_name',
            'department':           'department',
            'position':             'position',
            'hire date':            'hire_date',
            'hire_date':            'hire_date',
            'date regularized':     'date_regularized',
            'date_regularized':     'date_regularized',
            # Contact / gov ID columns from their actual Excel
            'phone no.':            'contact_no',
            'phone no':             'contact_no',
            'contact no':           'contact_no',
            'contact_no':           'contact_no',
            'address':              'address_1',
            'address 1':            'address_1',
            'address_1':            'address_1',
            'perosn to notify in':  'sos_contact_person',   # sic — matches their typo
            'person to notify in':  'sos_contact_person',
            'sos_contact_person':   'sos_contact_person',
            'phone #':              'sos_contact_no',
            'sos_contact_no':       'sos_contact_no',
            'tin':                  'tin_no',
            'tin no':               'tin_no',
            'tin_no':               'tin_no',
            'sss':                  'sss_no',
            'sss no':               'sss_no',
            'sss_no':               'sss_no',
            'philhealth':           'philhealth_no',
            'philhealth no':        'philhealth_no',
            'philhealth_no':        'philhealth_no',
            'pag-ibig':             'pagibig_no',
            'pagibig':              'pagibig_no',
            'pagibig_no':           'pagibig_no',
        }

        raw_headers = [
            str(c.value).strip().lower() if c.value is not None else ''
            for c in ws[1]
        ]
        col_map = {}
        for idx, h in enumerate(raw_headers):
            canonical = HEADER_ALIASES.get(h.strip())
            if canonical and canonical not in col_map:   # first match wins
                col_map[canonical] = idx

        missing_cols = [c for c in ['employee_id', 'first_name', 'last_name'] if c not in col_map]
        if missing_cols:
            return JsonResponse({'error': f'Missing required columns: {", ".join(missing_cols)}'}, status=400)

        departments = {d.name.strip().lower(): d for d in Department.objects.filter(is_active=True)}
        positions   = {p.title.strip().lower(): p for p in Position.objects.filter(is_active=True)}
        existing_employees = {
            emp.employee_id: f"{emp.first_name} {emp.last_name}"
            for emp in Employee.objects.only('employee_id', 'first_name', 'last_name')
        }
        existing_ids = set(existing_employees.keys())
        existing_names = {
            f"{emp.first_name.strip().lower()} {emp.last_name.strip().lower()}": emp.employee_id
            for emp in Employee.objects.only('employee_id', 'first_name', 'last_name')
        }

        def get_str(row, col):
            idx = col_map.get(col)
            if idx is None or idx >= len(row):
                return ''
            val = row[idx]
            return str(val).strip() if val is not None else ''

        def get_raw(row, col):
            idx = col_map.get(col)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        def parse_date(raw):
            if raw is None:
                return None, None
            from datetime import datetime as _dt, date as _d
            if isinstance(raw, _dt):
                return raw.strftime('%Y-%m-%d'), None
            if isinstance(raw, _d):
                return raw.strftime('%Y-%m-%d'), None
            s = str(raw).strip()
            if not s or s.lower() in ('none', 'nan', ''):
                return None, None
            for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%d/%m/%Y', '%m-%d-%Y', '%Y/%m/%d'):
                try:
                    return _dt.strptime(s, fmt).strftime('%Y-%m-%d'), None
                except ValueError:
                    pass
            return None, f'"{s}" could not be read as a date — will be left blank.'

        rows     = []
        seen_ids = set()

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(cell is not None for cell in row):
                continue  # skip fully blank rows

            first_name  = get_str(row, 'first_name')
            last_name   = get_str(row, 'last_name')
            middle_name_raw = get_str(row, 'middle_name')
            middle_name = '' if middle_name_raw.upper() in ('NM', 'N/A', 'N.A.', '-', 'NONE') else middle_name_raw
            suffix_name = get_str(row, 'suffix_name')
            emp_id_raw  = get_str(row, 'employee_id')
            dept_name   = get_str(row, 'department')
            pos_name    = get_str(row, 'position')

            issues = []
            status = 'valid'

            # ── Required field checks ──
            if not first_name:
                issues.append('First name is required.')
                status = 'error'
            if not last_name:
                issues.append('Last name is required.')
                status = 'error'

            emp_id = None
            if not emp_id_raw:
                issues.append('Employee ID is required.')
                status = 'error'
            else:
                try:
                    emp_id = int(float(emp_id_raw))
                except (ValueError, TypeError):
                    issues.append(f'Employee ID "{emp_id_raw}" is not a valid number.')
                    status = 'error'

            if emp_id is not None:
                if emp_id in existing_ids:
                    existing_name = existing_employees.get(emp_id, 'Unknown')
                    issues.append(f'Employee ID {emp_id} already exists in the system ({existing_name}). This row will be skipped.')
                    status = 'error'
                elif emp_id in seen_ids:
                    issues.append(f'Employee ID {emp_id} is duplicated within this file.')
                    status = 'error'
                else:
                    seen_ids.add(emp_id)

            # ── Same name, different ID check ──
            if emp_id is not None and status != 'error' and first_name and last_name:
                name_key = f"{first_name.strip().lower()} {last_name.strip().lower()}"
                matched_id = existing_names.get(name_key)
                if matched_id:
                    issues.append(f'An employee named {first_name} {last_name} already exists in the system (ID: {matched_id}). Please verify this is not a duplicate.')
                    if status == 'valid':
                        status = 'warning'

            # ── Department lookup ──
            dept_id = None
            if dept_name:
                dept = departments.get(dept_name.strip().lower())
                if dept:
                    dept_id = dept.id
                else:
                    issues.append(f'Department "{dept_name}" not found — will be left blank.')
                    if status == 'valid':
                        status = 'warning'

            # ── Position lookup ──
            pos_id = None
            if pos_name:
                pos = positions.get(pos_name.strip().lower())
                if pos:
                    pos_id = pos.id
                else:
                    issues.append(f'Position "{pos_name}" not found — will be left blank.')
                    if status == 'valid':
                        status = 'warning'

            # ── Date fields ──
            hire_date, hire_err = parse_date(get_raw(row, 'hire_date'))
            if hire_err:
                issues.append(hire_err)
                if status == 'valid':
                    status = 'warning'

            date_reg, reg_err = parse_date(get_raw(row, 'date_regularized'))
            if reg_err:
                issues.append(reg_err)
                if status == 'valid':
                    status = 'warning'

            rows.append({
                'row':               row_idx,
                'employee_id':       emp_id,
                'first_name':        first_name,
                'middle_name':       middle_name,
                'last_name':         last_name,
                'suffix_name':       suffix_name,
                'department_name':   dept_name,
                'department_id':     dept_id,
                'position_name':     pos_name,
                'position_id':       pos_id,
                'hire_date':         hire_date,
                'date_regularized':  date_reg,
                # Contact
                'contact_no':        get_str(row, 'contact_no'),
                'address_1':         get_str(row, 'address_1'),
                'sos_contact_person': get_str(row, 'sos_contact_person'),
                'sos_contact_no':    get_str(row, 'sos_contact_no'),
                # Gov IDs
                'tin_no':            get_str(row, 'tin_no'),
                'sss_no':            get_str(row, 'sss_no'),
                'philhealth_no':     get_str(row, 'philhealth_no'),
                'pagibig_no':        get_str(row, 'pagibig_no'),
                'status':            status,
                'issues':            issues,
            })

        summary = {
            'valid':   sum(1 for r in rows if r['status'] == 'valid'),
            'warning': sum(1 for r in rows if r['status'] == 'warning'),
            'error':   sum(1 for r in rows if r['status'] == 'error'),
        }
        return JsonResponse({'rows': rows, 'summary': summary})


@method_decorator([permission_required('employees', 'create')], name='dispatch')
class EmployeeBatchConfirmView(LoginRequiredMixin, View):

    def post(self, request):
        try:
            data = json.loads(request.body)
            rows = data.get('rows', [])
        except (json.JSONDecodeError, KeyError):
            return JsonResponse({'error': 'Invalid request body.'}, status=400)

        insertable = [r for r in rows if r.get('status') in ('valid', 'warning')]
        if not insertable:
            return JsonResponse({'error': 'No importable rows found.'}, status=400)

        created = skipped = 0
        errors  = []

        for r in insertable:
            try:
                emp_id = r.get('employee_id')
                if not emp_id or Employee.objects.filter(employee_id=emp_id).exists():
                    skipped += 1
                    continue

                dept = Department.objects.get(pk=r['department_id']) if r.get('department_id') else None
                pos  = Position.objects.get(pk=r['position_id'])     if r.get('position_id')  else None

                emp = Employee.objects.create(
                    employee_id      = emp_id,
                    first_name       = r.get('first_name', ''),
                    middle_name      = r.get('middle_name', ''),
                    last_name        = r.get('last_name', ''),
                    suffix_name      = r.get('suffix_name', ''),
                    department       = dept,
                    position         = pos,
                    hire_date        = r.get('hire_date') or None,
                    date_regularized = r.get('date_regularized') or None,
                )
                # Save contact info if any field is present
                contact_fields = ['contact_no', 'address_1', 'sos_contact_person', 'sos_contact_no']
                if any(r.get(f) for f in contact_fields):
                    EmployeeContact.objects.create(
                        employee           = emp,
                        contact_no         = r.get('contact_no', ''),
                        address_1          = r.get('address_1', ''),
                        sos_contact_person = r.get('sos_contact_person', ''),
                        sos_contact_no     = r.get('sos_contact_no', ''),
                    )
                # Save gov IDs if any field is present
                gov_fields = ['tin_no', 'sss_no', 'philhealth_no', 'pagibig_no']
                if any(r.get(f) for f in gov_fields):
                    EmployeeGovernmentID.objects.create(
                        employee      = emp,
                        tin_no        = r.get('tin_no', ''),
                        sss_no        = r.get('sss_no', ''),
                        philhealth_no = r.get('philhealth_no', ''),
                        pagibig_no    = r.get('pagibig_no', ''),
                    )
                created += 1
                
            except Exception as e:
                skipped += 1
                errors.append(f'Row {r.get("row", "?")}: {e}')

        return JsonResponse({'success': True, 'created': created, 'skipped': skipped, 'errors': errors})


@method_decorator([permission_required('employees', 'create')], name='dispatch')
class EmployeeBatchTemplateView(LoginRequiredMixin, View):

    def get(self, request):
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Employees'

        headers = ['employee_id', 'first_name', 'middle_name', 'last_name',
                   'suffix_name', 'department', 'position', 'hire_date', 'date_regularized']
        ws.append(headers)
        ws.append([10001, 'Juan', 'Santos', 'Dela Cruz', 'Jr.', 'Operations', 'Staff', '2024-01-15', '2024-07-15'])

        header_fill = PatternFill(start_color='1C4E80', end_color='1C4E80', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        for i, w in enumerate([12, 15, 15, 15, 10, 22, 22, 14, 18], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="employee_batch_template.xlsx"'
        return response

